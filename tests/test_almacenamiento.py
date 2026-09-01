"""Pruebas del almacenamiento: subida de fotos, compresión y respaldos.

El respaldo mensual BORRA datos de la base después de exportarlos, así que un
fallo silencioso aquí es pérdida de información irrecuperable.
"""
import io
import os
import shutil
from contextlib import contextmanager

import pytest
from PIL import Image, ImageDraw

from app.models.accesos import Acceso
from app.utils.imagenes import (
    extension_permitida,
    nombre_foto,
    procesar_foto,
    tiene_un_solo_rostro,
)

try:
    import cv2  # noqa: F401
    HAY_OPENCV = True
except ImportError:
    HAY_OPENCV = False


@contextmanager
def proteger_archivo_real(ruta, tmp_path):
    """Aísla una ruta dentro de una carpeta real del proyecto.

    Estas pruebas escriben en carpetas que la aplicación usa de verdad
    (`app/static/uploads/profiles/`, `app/respaldos_mensuales/`). Si ya
    existiera ahí un archivo real con el mismo nombre (una foto de perfil o un
    respaldo de producción), la prueba lo sobrescribiría y luego lo borraría.
    Este contexto lo aparta antes de la prueba y lo restaura al final, y de
    paso garantiza que lo que la prueba cree en esa ruta se elimina siempre,
    incluso si la prueba falla.
    """
    copia = None
    if os.path.isfile(ruta):
        copia = str(tmp_path / ('apartado_' + os.path.basename(ruta)))
        shutil.move(ruta, copia)
    try:
        yield
    finally:
        if os.path.isfile(ruta):
            os.remove(ruta)
        if copia:
            shutil.move(copia, ruta)


CARPETA_FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')


def foto_real(nombre):
    """Devuelve la ruta de una foto de referencia, o salta la prueba si falta.

    Las pruebas de aceptación necesitan una fotografía de verdad: el detector
    (YuNet) distingue fotos de ilustraciones a propósito, así que una cara
    dibujada con elipses la rechaza, y hace bien. Y una foto real no puede
    versionarse aquí porque es un dato biométrico de una persona identificable.

    Ver `tests/fixtures/README.md` para activarlas.
    """
    ruta = os.path.join(CARPETA_FIXTURES, nombre)
    if not os.path.isfile(ruta):
        pytest.skip(f"falta tests/fixtures/{nombre} (ver el README de esa carpeta)")
    return ruta


def imagen_dibujada(ancho=1600, alto=1600):
    """Imagen sintética con una cara dibujada.

    Sirve para probar el PROCESADO (reescalado, compresión, EXIF), que no
    necesita reconocer nada. No sirve para probar la validación facial: el
    detector la rechaza por ser un dibujo, que es justo lo que debe hacer.
    """
    img = Image.new('RGB', (ancho, alto), (235, 225, 215))
    d = ImageDraw.Draw(img)
    cx, cy, r = ancho // 2, alto // 2, min(ancho, alto) // 4
    d.ellipse([cx - r, cy - int(r * 1.25), cx + r, cy + int(r * 1.25)], fill=(228, 205, 185))
    oj = r // 4
    d.ellipse([cx - r // 2 - oj, cy - r // 2, cx - r // 2 + oj, cy - r // 2 + oj], fill=(40, 30, 25))
    d.ellipse([cx + r // 2 - oj, cy - r // 2, cx + r // 2 + oj, cy - r // 2 + oj], fill=(40, 30, 25))
    d.polygon([(cx, cy - r // 8), (cx - r // 8, cy + r // 4), (cx + r // 8, cy + r // 4)],
              fill=(200, 175, 155))
    d.arc([cx - r // 2, cy + r // 3, cx + r // 2, cy + r], 0, 180,
          fill=(120, 70, 60), width=max(3, r // 20))
    buf = io.BytesIO()
    img.save(buf, format='JPEG', quality=95)
    buf.seek(0)
    return buf


class TestExtensiones:
    @pytest.mark.parametrize('nombre', ['foto.jpg', 'FOTO.JPEG', 'a.png', 'b.webp'])
    def test_permitidas(self, nombre):
        assert extension_permitida(nombre) is True

    @pytest.mark.parametrize('nombre', ['x.svg', 'x.php', 'x.exe', 'sin_extension', '', None])
    def test_rechazadas(self, nombre):
        # SVG queda fuera a propósito: admite <script> y se sirve desde el
        # propio dominio, así que sería XSS almacenado.
        assert extension_permitida(nombre) is False

    def test_nombre_fijo_por_usuario(self):
        # Un nombre fijo hace que cada foto nueva reemplace a la anterior, en
        # vez de acumular imágenes huérfanas en el disco para siempre.
        assert nombre_foto(7) == 'user_7.jpg'


class TestProcesado:
    def test_reescala_comprime_y_convierte(self, tmp_path):
        destino = str(tmp_path / 'salida.jpg')
        origen = imagen_dibujada(2400, 2400)
        peso_original = len(origen.getvalue())
        origen.seek(0)

        assert procesar_foto(origen, destino) is True
        assert os.path.getsize(destino) < peso_original / 5

        with Image.open(destino) as im:
            assert max(im.size) <= 512
            assert im.format == 'JPEG'

    def test_elimina_los_metadatos_exif(self, tmp_path):
        # Las fotos de celular llevan coordenadas GPS en el EXIF: guardarlas
        # sería recoger un dato de ubicación que nadie autorizó.
        origen = tmp_path / 'con_exif.jpg'
        img = Image.new('RGB', (800, 800), (200, 180, 160))
        exif = img.getexif()
        exif[271] = 'FabricanteDePrueba'
        exif[272] = 'ModeloDePrueba'
        img.save(str(origen), exif=exif)
        assert Image.open(str(origen)).getexif(), 'la imagen de partida debe traer EXIF'

        destino = str(tmp_path / 'limpia.jpg')
        assert procesar_foto(str(origen), destino) is True
        with Image.open(destino) as im:
            assert not im.getexif()

    def test_aplana_la_transparencia_sobre_blanco(self, tmp_path):
        # Sin esto, un PNG con alfa sale con el fondo negro al pasar a JPEG.
        origen = tmp_path / 'transparente.png'
        Image.new('RGBA', (400, 400), (0, 0, 0, 0)).save(str(origen))
        destino = str(tmp_path / 'salida.jpg')
        assert procesar_foto(str(origen), destino) is True
        with Image.open(destino) as im:
            assert im.convert('RGB').getpixel((5, 5)) == (255, 255, 255)

    def test_rechaza_lo_que_no_es_una_imagen(self, tmp_path):
        # Abrir el archivo con Pillow es también la validación: un .jpg que en
        # realidad es un script falla aquí y nunca llega al disco.
        falso = io.BytesIO(b'<?php system($_GET["c"]); ?>')
        assert procesar_foto(falso, str(tmp_path / 'x.jpg')) is False

    def test_crea_la_carpeta_si_no_existe(self, tmp_path):
        destino = str(tmp_path / 'sub' / 'carpeta' / 'nueva.jpg')
        assert procesar_foto(imagen_dibujada(600, 600), destino) is True
        assert os.path.isfile(destino)


@pytest.mark.skipif(not HAY_OPENCV, reason="opencv no está instalado")
class TestValidacionFacial:
    """La foto de perfil se le muestra al celador para que confirme la
    identidad de quien entra, así que tiene que ser un retrato utilizable."""

    def test_acepta_un_retrato_real(self, tmp_path):
        ruta = str(tmp_path / 'retrato.jpg')
        procesar_foto(foto_real('rostro.jpg'), ruta)
        valida, mensaje = tiene_un_solo_rostro(ruta)
        assert valida is True, f"rechazó un retrato real: {mensaje}"

    def test_acepta_a_quien_lleva_gafas(self, tmp_path):
        # Requisito explícito del sistema. Si esto empezara a fallar, esa
        # persona se queda sin carnet digital y sin poder entrar al centro.
        ruta = str(tmp_path / 'gafas.jpg')
        procesar_foto(foto_real('rostro_gafas.jpg'), ruta)
        valida, mensaje = tiene_un_solo_rostro(ruta)
        assert valida is True, f"rechazó a alguien con gafas: {mensaje}"

    def test_acepta_el_retrato_en_distintas_condiciones(self, tmp_path):
        """Una foto tomada con el móvil de un aprendiz no llega perfecta:
        puede venir oscura, movida o girada. Debe aceptarse igual."""
        from PIL import ImageEnhance, ImageFilter

        original = Image.open(foto_real('rostro.jpg')).convert('RGB')
        variantes = {
            'oscura': ImageEnhance.Brightness(original).enhance(0.6),
            'clara': ImageEnhance.Brightness(original).enhance(1.4),
            'desenfocada': original.filter(ImageFilter.GaussianBlur(1.5)),
            'girada': original.rotate(6, expand=True, fillcolor=(255, 255, 255)),
            'pequena': original.resize((original.width // 3, original.height // 3)),
        }
        fallos = []
        for nombre, imagen in variantes.items():
            origen = tmp_path / f'{nombre}.jpg'
            imagen.save(str(origen), quality=92)
            destino = str(tmp_path / f'proc_{nombre}.jpg')
            procesar_foto(str(origen), destino)
            valida, mensaje = tiene_un_solo_rostro(destino)
            if not valida:
                fallos.append(f"{nombre}: {mensaje}")
        assert not fallos, "variantes rechazadas -> " + " | ".join(fallos)

    def test_rechaza_fotos_con_mas_de_una_persona(self, tmp_path):
        # Si salen dos personas, el celador no sabe a cuál corresponde el carnet.
        ruta = str(tmp_path / 'dos.jpg')
        procesar_foto(foto_real('rostro_dos_personas.jpg'), ruta)
        valida, mensaje = tiene_un_solo_rostro(ruta)
        assert valida is False
        assert 'más de una persona' in mensaje

    def test_rechaza_un_rostro_diminuto(self, tmp_path):
        """Una foto de cuerpo entero de lejos deja la cara con muy pocos
        píxeles: en el escáner no se distingue de nadie."""
        cara = Image.open(foto_real('rostro.jpg')).convert('RGB')
        cara.thumbnail((90, 90))
        escena = Image.new('RGB', (1400, 1400), (245, 245, 240))
        escena.paste(cara, (650, 650))
        origen = tmp_path / 'lejos.jpg'
        escena.save(str(origen), quality=92)
        destino = str(tmp_path / 'proc_lejos.jpg')
        procesar_foto(str(origen), destino)
        valida, _ = tiene_un_solo_rostro(destino)
        assert valida is False

    @pytest.mark.parametrize('color', [(10, 90, 200), (255, 255, 255), (0, 0, 0)])
    def test_rechaza_imagenes_sin_rostro(self, color, tmp_path):
        ruta = str(tmp_path / 'plana.jpg')
        buf = io.BytesIO()
        Image.new('RGB', (700, 700), color).save(buf, format='JPEG')
        buf.seek(0)
        procesar_foto(buf, ruta)
        valida, _ = tiene_un_solo_rostro(ruta)
        assert valida is False

    def test_rechaza_dibujos_y_capturas(self, tmp_path):
        """Alguien podría intentar poner un dibujo o una captura de pantalla
        como foto de perfil: entonces la verificación en portería no sirve."""
        casos = {}

        # Cara dibujada
        origen = tmp_path / 'dibujo.jpg'
        with open(str(origen), 'wb') as f:
            f.write(imagen_dibujada(900, 900).getvalue())
        casos['dibujo'] = str(origen)

        # Captura de pantalla simulada: bloques de texto sobre fondo claro
        captura = Image.new('RGB', (1000, 700), (250, 250, 250))
        d = ImageDraw.Draw(captura)
        d.rectangle([0, 0, 1000, 60], fill=(60, 90, 160))
        for i in range(12):
            d.rectangle([40, 100 + i * 45, 40 + (i % 5 + 3) * 110, 128 + i * 45],
                        fill=(210, 210, 215))
        ruta_captura = tmp_path / 'captura.jpg'
        captura.save(str(ruta_captura), quality=92)
        casos['captura'] = str(ruta_captura)

        # Icono: formas planas de colores
        icono = Image.new('RGB', (800, 800), (255, 255, 255))
        d = ImageDraw.Draw(icono)
        d.ellipse([150, 150, 650, 650], fill=(240, 160, 40))
        d.rectangle([300, 300, 500, 500], fill=(60, 60, 70))
        ruta_icono = tmp_path / 'icono.jpg'
        icono.save(str(ruta_icono), quality=92)
        casos['icono'] = str(ruta_icono)

        aceptados = []
        for nombre, origen_ruta in casos.items():
            destino = str(tmp_path / f'proc_{nombre}.jpg')
            procesar_foto(origen_ruta, destino)
            valida, _ = tiene_un_solo_rostro(destino)
            if valida:
                aceptados.append(nombre)
        assert not aceptados, f"aceptó como retrato: {aceptados}"

    def test_no_bloquea_el_sistema_si_falta_el_modelo(self, tmp_path, monkeypatch):
        """La validación es un apoyo, no un guardián: si el modelo no está,
        debe dejar pasar la foto en vez de impedir que nadie complete su
        perfil."""
        from app.utils import imagenes

        monkeypatch.setattr(imagenes, 'RUTA_MODELO_YUNET',
                            str(tmp_path / 'no_existe.onnx'))
        ruta = str(tmp_path / 'cualquiera.jpg')
        procesar_foto(imagen_dibujada(600, 600), ruta)
        assert imagenes.tiene_un_solo_rostro(ruta) == (True, None)


class TestSubidaCompleta:
    def _entrar(self, client, usuario):
        return client.post('/auth/login',
                           data={'correo': usuario.correo, 'password': 'Segura2026'})

    @pytest.mark.skipif(not HAY_OPENCV, reason="opencv no está instalado")
    def test_la_foto_se_guarda_comprimida_y_con_nombre_fijo(self, client, app,
                                                            crear_usuario, db,
                                                            tmp_path):
        usuario = crear_usuario(correo='ana@sena.edu.co', cargo='Aprendiz',
                                perfil_completo=False)
        self._entrar(client, usuario)

        carpeta = os.path.join(app.root_path, 'static', 'uploads', 'profiles')
        destino = os.path.join(carpeta, f'user_{usuario.id}.jpg')
        with proteger_archivo_real(destino, tmp_path):
            r = client.post('/usuarios/update_profile',
                            data={'foto': (open(foto_real('rostro.jpg'), 'rb'),
                                       'enorme.jpg')},
                            content_type='multipart/form-data',
                            headers={'X-Requested-With': 'XMLHttpRequest'})
            assert r.status_code == 200

            assert os.path.isfile(destino)
            db.session.refresh(usuario)
            assert usuario.foto == f'user_{usuario.id}.jpg'
            # Un temporal olvidado iría acumulando basura en cada subida.
            assert not [f for f in os.listdir(carpeta) if f.startswith('.tmp_')]

    def test_un_archivo_invalido_no_borra_la_foto_anterior(self, client, app,
                                                           crear_usuario, db,
                                                           tmp_path):
        # El procesado escribe primero a un temporal justo para esto.
        usuario = crear_usuario(correo='ana@sena.edu.co', cargo='Aprendiz')
        carpeta = os.path.join(app.root_path, 'static', 'uploads', 'profiles')
        os.makedirs(carpeta, exist_ok=True)
        anterior = os.path.join(carpeta, f'user_{usuario.id}.jpg')
        with proteger_archivo_real(anterior, tmp_path):
            with open(anterior, 'wb') as f:
                f.write(b'foto anterior')
            usuario.foto = f'user_{usuario.id}.jpg'
            db.session.commit()

            self._entrar(client, usuario)
            r = client.post('/usuarios/update_profile',
                            data={'foto': (io.BytesIO(b'no soy una imagen'), 'x.jpg')},
                            content_type='multipart/form-data',
                            headers={'X-Requested-With': 'XMLHttpRequest'})
            assert r.status_code == 400
            assert os.path.isfile(anterior)
            with open(anterior, 'rb') as f:
                assert f.read() == b'foto anterior'
            assert not [f for f in os.listdir(carpeta) if f.startswith('.tmp_')]


class TestRespaldoMensual:
    def test_no_mezcla_accesos_de_distintas_entidades(self, app, db, crear_usuario,
                                                      tmp_path):
        """referencia_id es polimórfico: sin filtrar por tipo_referencia, el
        acceso del visitante 7 se unía al usuario 7, se archivaba con la cédula
        de esa persona y después se borraba de la base."""
        from datetime import timedelta

        from app.utils import get_colombia_time
        from app.utils.respaldos import ejecutar_respaldo_mensual

        usuario = crear_usuario(correo='ana@sena.edu.co', documento='12345')
        mes_pasado = get_colombia_time().replace(day=1) - timedelta(days=5)
        db.session.add_all([
            Acceso(punto_id=1, referencia_id=usuario.id, tipo_referencia='Usuario',
                   tipo='Entrada', fecha=mes_pasado),
            Acceso(punto_id=1, referencia_id=usuario.id, tipo_referencia='Visitante',
                   tipo='Entrada', fecha=mes_pasado),
        ])
        db.session.commit()

        # Nombre exacto que va a generar el respaldo, no "el primer .xlsx que
        # haya en la carpeta": la carpeta es real y podría contener respaldos
        # de verdad que la prueba no debe tocar (ni inspeccionar ni borrar).
        ultimo_dia_mes_anterior = get_colombia_time().replace(day=1) - timedelta(days=1)
        carpeta = os.path.join(app.root_path, 'respaldos_mensuales')
        ruta = os.path.join(
            carpeta,
            f"Respaldo_Sistema_{ultimo_dia_mes_anterior.strftime('%Y-%m')}.xlsx")

        with proteger_archivo_real(ruta, tmp_path):
            ejecutar_respaldo_mensual()
            assert os.path.isfile(ruta), 'el respaldo debió generar un archivo'

            from openpyxl import load_workbook
            wb = load_workbook(ruta)
            assert 'Accesos No Usuarios' in wb.sheetnames, (
                'los accesos de visitantes y vehículos también deben archivarse; '
                'antes quedaban fuera y nunca se purgaban')

            filas = list(wb['Historial Accesos'].iter_rows(min_row=2, values_only=True))
            assert len(filas) == 1, 'solo el acceso del usuario va en esa hoja'

            # Todo lo archivado se purga: no debe quedar nada del mes anterior.
            assert Acceso.query.filter(Acceso.fecha < get_colombia_time().replace(day=1)).count() == 0


@pytest.mark.skipif(not HAY_OPENCV, reason="opencv no está instalado")
class TestRostroTapado:
    """El rostro debe verse despejado: con mascarilla o gorra calada, el celador
    no puede confirmar que quien está en la puerta es la persona de la foto.

    Las oclusiones se colocan usando los puntos faciales que devuelve el
    detector, no en fracciones fijas de la imagen: si se ponen "a ojo" pueden
    caer sobre el cuello o el fondo y la prueba pasa sin comprobar nada.
    """

    def _con_oclusion(self, ruta_origen, destino, modo, color):
        import cv2
        import numpy as np
        from PIL import Image, ImageDraw

        from app.utils.imagenes import RUTA_MODELO_YUNET

        procesar_foto(ruta_origen, destino)
        imagen = cv2.imdecode(np.fromfile(destino, dtype=np.uint8), cv2.IMREAD_COLOR)
        alto, ancho = imagen.shape[:2]
        detector = cv2.FaceDetectorYN_create(RUTA_MODELO_YUNET, '',
                                             (ancho, alto), 0.6, 0.3, 5000)
        _, rostros = detector.detect(imagen)
        assert rostros is not None, 'la foto de referencia debe tener un rostro'

        cara = rostros[0]
        x, y, ancho_r, alto_r = [int(v) for v in cara[:4]]
        puntos = [(int(cara[4 + i * 2]), int(cara[5 + i * 2])) for i in range(5)]
        nariz = puntos[2]
        ojos_y = min(puntos[0][1], puntos[1][1])

        img = Image.open(destino).convert('RGB')
        dibujo = ImageDraw.Draw(img)
        if modo == 'boca':
            dibujo.rectangle([x - 2, nariz[1] - int(alto_r * 0.06),
                              x + ancho_r + 2, y + alto_r + int(alto_r * 0.10)],
                             fill=color)
        else:
            dibujo.rectangle([x - int(ancho_r * 0.15), max(0, y - int(alto_r * 0.25)),
                              x + ancho_r + int(ancho_r * 0.15),
                              ojos_y - int(alto_r * 0.12)], fill=color)
        img.save(destino, quality=93)

    @pytest.mark.parametrize('etiqueta,modo,color', [
        ('mascarilla azul', 'boca', (120, 150, 190)),
        ('mascarilla negra', 'boca', (30, 30, 34)),
        # Del color de la piel: el caso que más cuesta, porque no se separa por color.
        ('mascarilla color piel', 'boca', (226, 200, 178)),
        ('bufanda', 'boca', (95, 70, 60)),
        ('gorra calada', 'frente', (40, 50, 90)),
    ])
    def test_rechaza_el_rostro_tapado(self, etiqueta, modo, color, tmp_path):
        destino = str(tmp_path / 'tapado.jpg')
        self._con_oclusion(foto_real('rostro.jpg'), destino, modo, color)
        valida, mensaje = tiene_un_solo_rostro(destino)
        assert valida is False, f'{etiqueta} debió rechazarse'
        assert 'tapad' in mensaje.lower() or 'gorra' in mensaje.lower()

    def test_las_gafas_siguen_aceptandose(self, tmp_path):
        # El filtro de oclusión no debe confundir unas gafas con algo que tapa
        # la cara: es un requisito explícito del sistema.
        destino = str(tmp_path / 'gafas.jpg')
        procesar_foto(foto_real('rostro_gafas.jpg'), destino)
        valida, mensaje = tiene_un_solo_rostro(destino)
        assert valida is True, f'rechazó a alguien con gafas: {mensaje}'
